import io
import json

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font
from streamlit_autorefresh import st_autorefresh

# ---------------------------
# НАСТРОЙКА СТРАНИЦЫ STREAMLIT
# ---------------------------
# Конфигурируем параметры страницы: заголовок, иконку и ширину макета
st.set_page_config(page_title="Аналитика Чат-Бота",
                   page_icon="🤖", layout="wide")

# ---------------------------
# CSS-ХАК ДЛЯ КНОПОК СКАЧИВАНИЯ
# ---------------------------
# Определяем стиль кнопок скачивания, чтобы задать фиксированные размеры и шрифт
st.markdown("""
<style>
button[data-testid="stDownloadButton"] {
    width: 120px !important;
    height: 35px !important;
    font-size: 12px;
    padding: 0 4px;
    margin-top: 4px;
}
</style>
""", unsafe_allow_html=True)

# ---------------------------
# АВТОМАТИЧЕСКАЯ ПЕРЕЗАГРУЗКА СТРАНИЦЫ
# ---------------------------
# Устанавливаем интервал автоперезагрузки (в данном случае 10 секунд)
time_interval = 10
st_autorefresh(interval=time_interval * 1000, key="data_refresh")


def make_headers_bold(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)


def set_manual_column_widths(ws):
    column_widths = {
        "A": 40, "B": 50, "C": 20, "D": 30, "E": 30, "F": 50, "G": 70
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width


def create_excel_file(normal, bad, unsure):
    wb = Workbook()
    headers = [
        "Вопрос", "Ответ AI", "Категория вопроса",
        "user_filters", "question_filters", "Очищенный контекст", "Все контексты"
    ]

    def make_headers_bold(ws):
        for cell in ws[1]:
            cell.font = Font(bold=True)

    def set_manual_column_widths(ws):
        column_widths = {
            "A": 40, "B": 50, "C": 20, "D": 30, "E": 30, "F": 50, "G": 70
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

    def fill_sheet(ws, data, title):
        ws.title = title
        ws.append(headers)
        make_headers_bold(ws)
        for case in data:
            ws.append([
                case.get("question", ""),
                case.get("answer", ""),
                case.get("question_category", ""),
                ", ".join(case.get("user_filters", [])),
                ", ".join(case.get("question_filters", [])),
                case.get("ground_truth", ""),
                "\n---\n".join(case.get("contexts", []))
            ])
        set_manual_column_widths(ws)

    ws1 = wb.active
    fill_sheet(ws1, normal, "Вопросы с ответом")
    fill_sheet(wb.create_sheet(), bad, "Вопросы без ответа")
    fill_sheet(wb.create_sheet(), unsure, "Частичный ответ")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ---------------------------
# ФУНКЦИЯ ЗАГРУЗКИ ДАННЫХ
# ---------------------------
def load_data(file_name: str):
    """
    Загружает JSON-файл и возвращает его содержимое.
    :param file_name: Имя файла с данными.
    :return: Данные в формате Python (список словарей).
    """
    with open(file_name, "r", encoding="utf-8") as file:
        return json.load(file)


# ---------------------------
# ФУНКЦИЯ ОБРАБОТКИ ДАННЫХ
# ---------------------------
def process_data(data):
    """
    Преобразует список данных в DataFrame и рассчитывает дополнительные метрики.
    Рассчитывается:
      - conflict_metric: показывает наличие конфликтов (например, если в истории чата более одного уточняющего вопроса и время ответа > 3 сек)
      - has_chat_history/has_contexts: флаги наличия истории чата или контекстов.
    Также происходит приведение времени ответа к числовому типу.
    :param data: Сырые данные (список словарей).
    :return: Обработанный DataFrame.
    """
    df = pd.DataFrame(data)

    # Если есть столбец chat_history, вычисляем метрики по истории чата
    if "chat_history" in df.columns:
        df["has_chat_history"] = df["chat_history"].apply(
            lambda x: len(x.get("old_questions", [])) > 0)
        df["conflict_metric"] = df.apply(
            lambda row: 1 if (len(row.get("chat_history", {}).get("old_questions", [])) > 1
                              and row["response_time"] > 3) else 0,
            axis=1
        )
    # Если нет chat_history, но есть contexts, делаем аналогичные вычисления
    elif "contexts" in df.columns:
        df["has_contexts"] = df["contexts"].apply(
            lambda x: len(x) > 0 if isinstance(x, list) else False)
        df["conflict_metric"] = df.apply(
            lambda row: 1 if (row["has_contexts"] and len(row.get("contexts", [])) > 1
                              and row["response_time"] > 3) else 0,
            axis=1
        )
    else:
        df["conflict_metric"] = 0

    # Приводим время ответа к числовому типу для последующих вычислений
    df["response_time"] = pd.to_numeric(df["response_time"], errors="coerce")
    return df


# ---------------------------
# ФУНКЦИЯ ДЛЯ СКАЧИВАНИЯ JSON ДАННЫХ
# ---------------------------
def download_json(data):
    """
    Преобразует данные в формат JSON и выводит кнопку для скачивания.
    :param data: Данные, которые нужно сохранить.
    """
    json_data = json.dumps(data, indent=4, ensure_ascii=False, default=str)
    st.download_button(
        label="📥 Скачать JSON",
        data=json_data,
        file_name="chatbot_logs.json",
        mime="application/json"
    )


# ---------------------------
# ВСПОМОГАТЕЛЬНАЯ ФУНКЦИЯ ДЛЯ ОТОБРАЖЕНИЯ ГРАФИКОВ
# ---------------------------
def show_plot_with_download_below(fig, filename: str):
    """
    Отображает график Plotly и добавляет под ним кнопку для скачивания в формате PNG.
    :param fig: Объект графика Plotly.
    :param filename: Имя файла для сохранения.
    """
    st.plotly_chart(fig, use_container_width=True)
    try:
        img_bytes = fig.to_image(format="png")
        st.download_button(
            label="Скачать график",
            data=img_bytes,
            file_name=f"{filename}.png",
            mime="image/png"
        )
    except Exception as e:
        st.error(f"Ошибка экспорта: {e}")


# ---------------------------
# КЛАСС ДЛЯ ПОСТРОЕНИЯ ГРАФИКОВ
# ---------------------------
class Plots:
    def __init__(self, data: pd.DataFrame):
        """
        Инициализирует объект для построения графиков.
        :param data: Обработанный DataFrame с данными.
        """
        self.data = data

    # 1. Построение пироговой диаграммы
    def plot_pie_chart(self, column: str, _unused_title: str):
        """
        Строит пироговую диаграмму для указанного столбца.
        :param column: Имя столбца для агрегации.
        :param _unused_title: Не используется (оставлено для совместимости).
        """
        if self.data.empty or column not in self.data.columns or self.data[column].dropna().empty:
            return st.info("Нет данных для построения графика")
        counts = self.data[column].value_counts()
        fig = px.pie(
            names=counts.index,
            values=counts.values,
            hole=0.4,  # создание полого пирога
            color_discrete_sequence=px.colors.sequential.RdBu
        )
        show_plot_with_download_below(fig, f"pie_{column}")

    # 2. Построение столбчатой диаграммы
    def plot_bar_chart(self, column: str, _unused_title: str, x_label: str, y_label: str):
        """
        Строит столбчатую диаграмму для указанного столбца.
        :param column: Имя столбца для агрегации.
        :param _unused_title: Не используется.
        :param x_label: Подпись оси X.
        :param y_label: Подпись оси Y.
        """
        if self.data.empty or column not in self.data.columns or self.data[column].dropna().empty:
            return st.info("Нет данных для построения графика")
        counts = self.data[column].value_counts()
        if counts.empty:
            return st.info("Нет данных для построения графика")
        fig = px.bar(
            x=counts.index,
            y=counts.values,
            labels={'x': x_label, 'y': y_label},
            text_auto=True,  # автоматическое отображение значений над столбцами
            color_discrete_sequence=px.colors.qualitative.Vivid
        )
        show_plot_with_download_below(fig, f"bar_{column}")

    # 3. Среднее время ответа по кампусам
    def plot_response_time_chart_with_campus(self):
        """
        Строит столбчатую диаграмму, показывающую среднее время ответа для каждого кампуса.
        """
        if self.data.empty or "campus" not in self.data.columns or "response_time" not in self.data.columns:
            return st.info("Нет данных для построения графика")
        group_data = self.data.groupby(
            "campus")["response_time"].mean().reset_index()
        if group_data.empty:
            return st.info("Нет данных для построения графика")
        fig = px.bar(
            group_data,
            x="campus",
            y="response_time",
            color="campus",
            text_auto=True,
            color_discrete_sequence=px.colors.qualitative.Set3
        )
        show_plot_with_download_below(fig, "resp_time_by_campus")

    # 4. Усреднение времени ответа по группам (например, по 10 запросов)
    def plot_averaged_response_time_chart(self, bin_size: int = 10):
        """
        Строит график, показывающий усредненное время ответа для групп запросов.
        :param bin_size: Количество запросов в одной группе.
        """
        if self.data.empty or "response_time" not in self.data.columns:
            return st.info("Нет данных для построения графика")
        df_copy = self.data.copy()
        df_copy["group"] = df_copy.index // bin_size
        grouped = df_copy.groupby(
            "group")["response_time"].mean().reset_index()
        fig = px.bar(
            grouped,
            x="group",
            y="response_time",
            labels={
                "group": f"Номер группы (по {bin_size} запросов)",
                "response_time": "Среднее время ответа (сек)"
            }
        )
        show_plot_with_download_below(fig, "resp_time_averaged")

    # 5. Пироговая диаграмма для доли уточняющих вопросов
    def plot_follow_up_pie_chart(self):
        """
        Строит пироговую диаграмму, показывающую процент запросов с уточнениями.
        """
        if self.data.empty:
            return st.info("Нет данных для построения графика")
        # Выбор столбца в зависимости от наличия истории чата или контекстов
        flag = "has_chat_history" if "has_chat_history" in self.data.columns else "has_contexts"
        if flag not in self.data.columns or self.data[flag].dropna().empty:
            return st.info("Нет данных для построения графика")
        avg_flag = self.data[flag].mean()
        fig = px.pie(
            names=["Без уточнений", "С уточнениями"],
            values=[1 - avg_flag, avg_flag],
            hole=0.3,
            color_discrete_sequence=px.colors.qualitative.Pastel
        )
        show_plot_with_download_below(fig, "follow_up_pie")

    # 6. Гейдж для отображения метрики конфликтного ответа
    def plot_conflict_metric(self):
        """
        Строит гейдж (индикатор) для отображения среднего значения conflict_metric (в процентах).
        """
        if self.data.empty or "conflict_metric" not in self.data.columns:
            return st.info("Нет данных для построения графика")
        conflict_rate = self.data["conflict_metric"].mean() * 100
        fig = go.Figure(go.Indicator(
            mode="gauge+number",
            value=conflict_rate,
            gauge={
                'axis': {'range': [0, 100]},
                'bar': {'color': "red"},
                'steps': [{'range': [0, 100], 'color': "lightcoral"}],
            }
        ))
        show_plot_with_download_below(fig, "conflict_metric")

    # 7. Среднее время ответа по категориям вопросов
    def plot_response_time_by_category(self):
        """
        Строит столбчатую диаграмму, показывающую среднее время ответа для каждой категории вопросов.
        """
        if self.data.empty or "question_category" not in self.data.columns or "response_time" not in self.data.columns:
            return st.info("Нет данных для построения графика")
        grouped = self.data.groupby("question_category")[
            "response_time"].mean().reset_index()
        if grouped.empty:
            return st.info("Нет данных для построения графика")
        fig = px.bar(
            grouped,
            x="question_category",
            y="response_time",
            color_discrete_sequence=px.colors.qualitative.Pastel
        )
        show_plot_with_download_below(fig, "resp_time_by_category")

    # 8. BoxPlot для времени ответа
    def plot_response_time_boxplot(self):
        """
        Строит boxplot для распределения времени ответа.
        """
        if self.data.empty or "response_time" not in self.data.columns:
            return st.info("Нет данных для построения графика")
        fig = px.box(
            self.data,
            y="response_time",
            color_discrete_sequence=["#FF6666"]
        )
        show_plot_with_download_below(fig, "resp_time_boxplot")

    # 9. Построение отдельных графиков для метрик качества
    def plot_quality_metrics_separate(self):
        """
        Строит отдельный столбчатый график для каждой из метрик качества:
        - context_recall
        - context_precision
        - answer_correctness_literal
        - answer_correctness_neural
        - Hallucination_metric
        Графики строятся в 5 колонках.
        """
        needed_cols = [
            "question_category",
            "context_recall",
            "context_precision",
            "answer_correctness_literal",
            "answer_correctness_neural",
            "Hallucination_metric"
        ]
        for c in needed_cols:
            if c not in self.data.columns:
                return st.info(f"Нет столбца '{c}' для построения метрик.")

        # Список метрик для построения
        metrics = [
            "context_recall",
            "context_precision",
            "answer_correctness_literal",
            "answer_correctness_neural",
            "Hallucination_metric"
        ]

        # Создаем 5 колонок для отображения графиков в одной строке
        cols = st.columns(5)
        for i, metric in enumerate(metrics):
            grouped = self.data.groupby("question_category")[
                metric].mean().reset_index()
            fig = px.bar(
                grouped,
                x="question_category",
                y=metric,
                labels={
                    "question_category": "Категория вопроса",
                    metric: "Среднее значение"
                },
                title=f"Метрика: {metric}"
            )
            with cols[i]:
                show_plot_with_download_below(fig, f"separate_{metric}")

    # 10. Сводный график для всех метрик качества на одном полотне
    def plot_quality_metrics_combined(self):
        """
        Строит один сводный столбчатый график, где все 5 метрик качества
        (context_recall, context_precision, answer_correctness_literal, answer_correctness_neural, Hallucination_metric)
        приводятся к диапазону [0, 100] и отображаются группой для каждой категории вопросов.
        """
        needed_cols = [
            "question_category",
            "context_recall",
            "context_precision",
            "answer_correctness_literal",
            "answer_correctness_neural",
            "Hallucination_metric"
        ]
        for c in needed_cols:
            if c not in self.data.columns:
                return st.info(f"Нет столбца '{c}' для построения метрик.")

        metrics = [
            "context_recall",
            "context_precision",
            "answer_correctness_literal",
            "answer_correctness_neural",
            "Hallucination_metric"
        ]

        # Группировка данных по категориям вопросов и вычисление средних значений для каждой метрики
        grouped = self.data.groupby("question_category")[
            metrics].mean().reset_index()

        # Масштабирование значений каждой метрики к диапазону [0, 100] для корректного сравнения
        for metric in metrics:
            max_val = grouped[metric].max()
            if max_val > 0:
                grouped[metric] = grouped[metric] / max_val * 100

        # Преобразование данных в длинный формат для построения группированного графика
        melted = grouped.melt(
            id_vars="question_category",
            value_vars=metrics,
            var_name="metric",
            value_name="mean_value"
        )

        fig = px.bar(
            melted,
            x="question_category",
            y="mean_value",
            color="metric",
            barmode="group",
            labels={
                "question_category": "Категория вопроса",
                "mean_value": "Среднее (0–100)",
                "metric": "Метрика"
            },
            title=""  # Заголовок не задаем, чтобы не дублировать надпись на странице
        )
        show_plot_with_download_below(fig, "combined_quality_metrics")


# ---------------------------
# ФУНКЦИЯ САЙДБАРА ДЛЯ ФИЛЬТРАЦИИ ДАННЫХ
# ---------------------------
def sidebar_layout(df: pd.DataFrame):
    """
    Отображает боковую панель с фильтрами для данных.
    Позволяет пользователю выбрать кампус, категорию вопроса и уровень образования.
    Возвращает отфильтрованный DataFrame.
    """
    st.sidebar.image(
        "https://github.com/X-D-R/hackathon_hse25/raw/main/logo.png",
        use_container_width=True
    )
    st.sidebar.title("Фильтры")

    campuses = df["campus"].dropna().unique(
    ).tolist() if "campus" in df.columns else []
    categories = df["question_category"].dropna().unique(
    ).tolist() if "question_category" in df.columns else []
    education_levels = df["education_level"].dropna().unique(
    ).tolist() if "education_level" in df.columns else []

    selected_campus = st.sidebar.multiselect(
        "Выберите кампус", campuses, default=campuses)
    selected_category = st.sidebar.multiselect(
        "Выберите категорию вопроса", categories, default=categories)
    selected_edu_level = st.sidebar.multiselect("Выберите уровень образования", education_levels,
                                                default=education_levels)

    filtered_df = df.copy()
    if "campus" in df.columns:
        filtered_df = filtered_df[filtered_df["campus"].isin(selected_campus)]
    if "question_category" in df.columns:
        filtered_df = filtered_df[filtered_df["question_category"].isin(
            selected_category)]
    if "education_level" in df.columns:
        filtered_df = filtered_df[filtered_df["education_level"].isin(
            selected_edu_level)]

    return filtered_df


# --- Excel: подготовка и экспорт ---
def split_by_answer_quality(data):
    normal, bad, unsure = [], [], []
    for case in data:
        answer = case.get("answer", "")
        if not answer.strip():
            bad.append(case)
        elif any(x in answer.lower() for x in ["возможно", "не уверен", "не могу сказать"]):
            unsure.append(case)
        else:
            normal.append(case)
    return normal, bad, unsure


# ---------------------------
# ГЛАВНАЯ ФУНКЦИЯ
# ---------------------------
def main():
    """
    Основная функция приложения.
    Загружает и обрабатывает данные, отображает фильтры, а затем строит различные графики:
      - Отдельные графики для метрик качества
      - Сводный график для метрик качества
      - Графики основных метрик (распределение по кампусам, уровням образования и т.д.)
      - Графики, связанные с временем ответа и дополнительными метриками
    """
    # Загрузка данных из обновленного JSON-файла
    data = load_data("output_last (1).json")
    df = process_data(data)
    filtered_df = sidebar_layout(df)
    if filtered_df.empty:
        st.info("Нет данных для отображения. Попробуйте изменить фильтры.")
        return

    graphs = Plots(filtered_df)

    # Заголовок приложения
    st.markdown("<h1 style='text-align: center;'>Мониторинг качества чат-бота</h1>",
                unsafe_allow_html=True)

    # Кнопка для скачивания отфильтрованных данных в формате JSON
    st.markdown("### Экспорт данных")

    col1, col2 = st.columns(2)
    with col1:
        download_json(filtered_df.to_dict(orient="records"))
    with col2:
        normal_q, bad_q, unsure_q = split_by_answer_quality(
            filtered_df.to_dict(orient="records"))
        excel_file = create_excel_file(normal_q, bad_q, unsure_q)

        st.download_button(
            label="📥 Скачать Excel",
            data=excel_file,
            file_name="chatbot_logs.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    # --- 1) Отдельные графики для метрик качества ---
    st.markdown("## Отдельные метрики качества")
    graphs.plot_quality_metrics_separate()

    # --- 2) Сводный график метрик качества ---
    st.markdown("## Сводный график метрик качества")
    graphs.plot_quality_metrics_combined()

    # --- 3) Основные метрики ---
    st.markdown("## Основные метрики")
    col1, col2, col3 = st.columns(3)
    with col1:
        st.subheader("Распределение запросов по кампусам")
        if "campus" in filtered_df.columns:
            graphs.plot_pie_chart("campus", "unused_title")
        else:
            st.info("Нет столбца 'campus'")
    with col2:
        st.subheader("Распределение по уровням образования")
        if "education_level" in filtered_df.columns:
            graphs.plot_pie_chart("education_level", "unused_title")
        else:
            st.info("Нет столбца 'education_level'")
    with col3:
        st.subheader("Частота уточняющих вопросов")
        graphs.plot_follow_up_pie_chart()

    # --- 4) Сравнение времени ответа по категориям ---
    st.markdown("## Сравнение времени ответа по категориям")
    graphs.plot_response_time_by_category()

    # --- 5) Среднее время ответа (по кампусам и группам) ---
    st.markdown("## Сравнения по времени ответа")
    col4, col5 = st.columns(2)
    with col4:
        st.subheader("Среднее время ответа по кампусам")
        graphs.plot_response_time_chart_with_campus()
    with col5:
        st.subheader("Усреднённое время ответа (по группам)")
        graphs.plot_averaged_response_time_chart(bin_size=10)

    # --- 6) Дополнительные графики ---
    st.markdown("## Дополнительные графики")
    st.subheader("Распределение времени ответа (BoxPlot)")
    graphs.plot_response_time_boxplot()
    st.subheader("Метрика конфликтного ответа")
    graphs.plot_conflict_metric()


if __name__ == "__main__":
    main()
